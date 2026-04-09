# src/services/signal_label_mapper.py
from __future__ import annotations
from pathlib import Path
import re
import openpyxl

_PREFIX_RE = re.compile(r"^\[.*?\]")   # [3_CO_PLC_A] 등 제거
_SPACE_RE  = re.compile(r"\s+")         # 연속 공백 → 단일 공백


def _normalize_tag(tag: str) -> str:
    """PLC prefix 제거 + 공백 정규화 + no-space 주소 보정."""
    tag = _PREFIX_RE.sub("", tag).strip()
    # G-5: DB420.DBW0 → DB420.DBW 0 (4OVEN xlsx 공백 없는 주소 보정)
    tag = re.sub(r"(DB\w+\.DB[WDXB])(\d)", r"\1 \2", tag)
    tag = _SPACE_RE.sub(" ", tag)
    return tag


class SignalLabelMapper:
    """OVEN xlsx에서 PLC 태그 주소 → (변수명, 설명) 매핑을 제공한다.

    xlsx 파일이 없거나 매핑이 없으면 None을 반환한다 (degraded gracefully).

    xlsx 구조 (각 시트의 컬럼):
        Col 0: PLC 태그 주소  예) DB420.DBW 2
        Col 1: 변수명          예) FR_HMI_W2
        Col 2: 데이터타입      예) INT
        Col 3: 설명            예) C/C STAND PIPE OPEN COUNT SET VALUE
    """

    def __init__(self, xlsx_paths: list[str | Path]):
        self._map: dict[str, tuple[str, str]] = {}  # {norm_tag: (var_name, description)}
        self._unmapped: list[str] = []               # 매핑 실패 태그 (런타임 기록용)
        for path in xlsx_paths:
            self._load_xlsx(Path(path))

    def _load_xlsx(self, path: Path) -> None:
        if not path.exists():
            return
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows(min_row=2, values_only=True):
                tag_raw = row[0]
                var_name = row[1] if len(row) > 1 else None
                description = row[3] if len(row) > 3 else None
                if not tag_raw or not isinstance(tag_raw, str):
                    continue
                norm = _normalize_tag(tag_raw)
                if norm and norm not in self._map:
                    self._map[norm] = (
                        str(var_name) if var_name else "",
                        str(description) if description else "",
                    )
        wb.close()

    def lookup(self, tag: str) -> tuple[str, str] | None:
        """정규화된 태그명으로 (변수명, 설명)을 반환한다. 없으면 None."""
        norm = _normalize_tag(tag)
        result = self._map.get(norm)
        if result is None:
            self._unmapped.append(tag)
        return result

    def label(self, tag: str, fallback: str | None = None) -> str:
        """변수명을 반환한다. 없으면 fallback 또는 원래 tag."""
        hit = self.lookup(tag)
        if hit and hit[0]:
            return hit[0]
        return fallback if fallback is not None else tag

    def describe(self, tag: str) -> str:
        """설명을 반환한다. 없으면 빈 문자열."""
        hit = self.lookup(tag)
        return hit[1] if hit else ""

    @property
    def size(self) -> int:
        return len(self._map)

    @property
    def unmapped_tags(self) -> list[str]:
        """런타임 중 매핑되지 않은 태그 목록 (중복 포함)."""
        return list(self._unmapped)
